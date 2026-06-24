from __future__ import annotations

import json
import mimetypes
import shutil
from dataclasses import dataclass
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any
from urllib.parse import unquote, urlparse

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
PUBLIC_DIR = ROOT / "public"
WORKBOOK_PATH = ROOT / "salvo equation .xlsx"
BACKUP_DIR = ROOT / "backups"

SALVO_VALUES = {"": 0, "blank": 0, "minimum": 1, "optimum": 2, "maximum": 3}


@dataclass(frozen=True)
class ForceConfig:
    key: str
    sheet: str
    label: str
    color: str
    start_row: int
    end_row: int
    total_row: int
    unit_total_end_row: int


FORCES = {
    "red": ForceConfig("red", "red force", "Red Force", "red", 5, 21, 22, 21),
    "blue": ForceConfig("blue", "blue force", "Blue Force", "blue", 5, 27, 28, 20),
}


def as_number(value: Any, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def clean_number(value: float) -> int | float:
    if abs(value - round(value)) < 1e-10:
        return int(round(value))
    return round(value, 6)


def read_force_sheet(workbook: Any, config: ForceConfig) -> dict[str, Any]:
    ws = workbook[config.sheet]
    salvo_mode = str(ws["C2"].value or "").strip().lower()
    salvo_factor = SALVO_VALUES.get(salvo_mode, 0)
    rows = []

    for row_number in range(config.start_row, config.end_row + 1):
        count = as_number(ws[f"D{row_number}"].value)
        missile_number = as_number(ws[f"E{row_number}"].value)
        missiles_total = count * missile_number
        effective_salvo = as_number(ws[f"G{row_number}"].value)
        asmd = as_number(ws[f"I{row_number}"].value)
        neutralize = as_number(ws[f"J{row_number}"].value)
        fire_power = missiles_total * salvo_factor * effective_salvo
        defense_power = count * asmd
        staying_power = count * neutralize

        rows.append(
            {
                "row": row_number,
                "countInUnitTotal": row_number <= config.unit_total_end_row,
                "displayNumber": ws[f"N{row_number}"].value,
                "unit": ws[f"C{row_number}"].value or "",
                "number": clean_number(count),
                "missileNumber": clean_number(missile_number),
                "missilesTotal": clean_number(missiles_total),
                "salvoSize": salvo_factor,
                "effectiveSalvo": clean_number(effective_salvo),
                "firePower": clean_number(fire_power),
                "asmdCapability": clean_number(asmd),
                "neutralizeHits": clean_number(neutralize),
                "defensePower": clean_number(defense_power),
                "stayingPower": clean_number(staying_power),
            }
        )

    totals = calculate_totals(rows)
    return {
        "key": config.key,
        "label": config.label,
        "color": config.color,
        "sheet": config.sheet,
        "salvoMode": salvo_mode if salvo_mode in SALVO_VALUES else "",
        "salvoFactor": salvo_factor,
        "rows": rows,
        "totals": totals,
    }


def calculate_totals(rows: list[dict[str, Any]]) -> dict[str, int | float]:
    return {
        "units": clean_number(
            sum(as_number(row.get("number")) for row in rows if row.get("countInUnitTotal", True))
        ),
        "firePower": clean_number(sum(as_number(row.get("firePower")) for row in rows)),
        "missilesTotal": clean_number(sum(as_number(row.get("missilesTotal")) for row in rows)),
        "defensePower": clean_number(sum(as_number(row.get("defensePower")) for row in rows)),
        "stayingPower": clean_number(sum(as_number(row.get("stayingPower")) for row in rows)),
    }


def calculate_combat(red: dict[str, Any], blue: dict[str, Any]) -> dict[str, Any]:
    red_totals = red["totals"]
    blue_totals = blue["totals"]
    blue_staying = as_number(blue_totals["stayingPower"])
    red_staying = as_number(red_totals["stayingPower"])

    blue_damage = (
        (as_number(red_totals["firePower"]) - as_number(blue_totals["defensePower"]))
        / blue_staying
        if blue_staying
        else 0
    )
    red_damage = (
        (as_number(blue_totals["firePower"]) - as_number(red_totals["defensePower"]))
        / red_staying
        if red_staying
        else 0
    )

    if abs(blue_damage - red_damage) < 1e-9:
        advantage = "balanced"
    elif blue_damage > red_damage:
        advantage = "red"
    else:
        advantage = "blue"

    return {
        "blueDamageByRed": round(blue_damage, 10),
        "redDamageByBlue": round(red_damage, 10),
        "advantage": advantage,
        "formula": {
            "blueDamageByRed": "(Red Fire Power - Blue Defense Power) / Blue Staying Power",
            "redDamageByBlue": "(Blue Fire Power - Red Defense Power) / Red Staying Power",
        },
    }


def load_state() -> dict[str, Any]:
    workbook = openpyxl.load_workbook(WORKBOOK_PATH, data_only=False)
    red = read_force_sheet(workbook, FORCES["red"])
    blue = read_force_sheet(workbook, FORCES["blue"])
    return {
        "workbook": {
            "path": str(WORKBOOK_PATH),
            "name": WORKBOOK_PATH.name,
            "lastModified": WORKBOOK_PATH.stat().st_mtime,
        },
        "forces": {"red": red, "blue": blue},
        "combat": calculate_combat(red, blue),
    }


def coerce_payload_row(row: dict[str, Any], fallback_row: int) -> dict[str, Any]:
    return {
        "row": int(row.get("row") or fallback_row),
        "unit": str(row.get("unit") or "").strip(),
        "number": as_number(row.get("number")),
        "missileNumber": as_number(row.get("missileNumber")),
        "effectiveSalvo": as_number(row.get("effectiveSalvo")),
        "asmdCapability": as_number(row.get("asmdCapability")),
        "neutralizeHits": as_number(row.get("neutralizeHits")),
    }


def write_force_sheet(workbook: Any, config: ForceConfig, force_payload: dict[str, Any]) -> None:
    ws = workbook[config.sheet]
    salvo_mode = str(force_payload.get("salvoMode") or "").strip().lower()
    if salvo_mode not in SALVO_VALUES:
        salvo_mode = ""
    ws["C2"] = salvo_mode

    rows_by_number = {
        int(row.get("row")): row for row in force_payload.get("rows", []) if row.get("row")
    }
    for row_number in range(config.start_row, config.end_row + 1):
        payload_row = coerce_payload_row(rows_by_number.get(row_number, {}), row_number)
        ws[f"C{row_number}"] = payload_row["unit"] or None
        ws[f"D{row_number}"] = payload_row["number"]
        ws[f"E{row_number}"] = payload_row["missileNumber"]
        ws[f"H{row_number}"] = f"=D{row_number}*E{row_number}*F{row_number}*G{row_number}"
        ws[f"G{row_number}"] = payload_row["effectiveSalvo"]
        ws[f"I{row_number}"] = payload_row["asmdCapability"]
        ws[f"J{row_number}"] = payload_row["neutralizeHits"]


def save_state(payload: dict[str, Any]) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(WORKBOOK_PATH, data_only=False)
    forces_payload = payload.get("forces", {})
    write_force_sheet(workbook, FORCES["red"], forces_payload.get("red", {}))
    write_force_sheet(workbook, FORCES["blue"], forces_payload.get("blue", {}))

    try:
        workbook.calculation.calcMode = "auto"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
    except AttributeError:
        pass

    BACKUP_DIR.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup_path = BACKUP_DIR / f"{WORKBOOK_PATH.stem}-{stamp}{WORKBOOK_PATH.suffix}"
    shutil.copy2(WORKBOOK_PATH, backup_path)
    workbook.save(WORKBOOK_PATH)
    return load_state() | {"backup": str(backup_path)}


class AppHandler(BaseHTTPRequestHandler):
    server_version = "SalvoWeb/1.0"

    def log_message(self, format: str, *args: Any) -> None:
        print(f"{self.address_string()} - {format % args}")

    def send_json(self, status: int, payload: dict[str, Any]) -> None:
        data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self) -> None:
        parsed = urlparse(self.path)
        if parsed.path == "/api/state":
            try:
                self.send_json(200, load_state())
            except Exception as exc:  # noqa: BLE001
                self.send_json(500, {"error": str(exc)})
            return

        if parsed.path == "/":
            relative = "index.html"
        else:
            relative = unquote(parsed.path.lstrip("/"))
        file_path = (PUBLIC_DIR / relative).resolve()

        if not str(file_path).startswith(str(PUBLIC_DIR.resolve())) or not file_path.exists():
            self.send_error(404)
            return

        content = file_path.read_bytes()
        mime_type, _ = mimetypes.guess_type(file_path.name)
        self.send_response(200)
        self.send_header("Content-Type", mime_type or "application/octet-stream")
        self.send_header("Content-Length", str(len(content)))
        self.end_headers()
        self.wfile.write(content)

    def do_POST(self) -> None:
        if urlparse(self.path).path != "/api/save":
            self.send_error(404)
            return

        try:
            length = int(self.headers.get("Content-Length") or 0)
            payload = json.loads(self.rfile.read(length).decode("utf-8"))
            self.send_json(200, save_state(payload))
        except PermissionError:
            self.send_json(
                423,
                {
                    "error": "Cannot write the Excel file. Close it in Excel and try Save again.",
                },
            )
        except Exception as exc:  # noqa: BLE001
            self.send_json(500, {"error": str(exc)})


def main() -> None:
    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError(f"Workbook not found: {WORKBOOK_PATH}")
    address = ("127.0.0.1", 8765)
    print(f"Salvo web app running at http://{address[0]}:{address[1]}")
    print(f"Using workbook: {WORKBOOK_PATH}")
    ThreadingHTTPServer(address, AppHandler).serve_forever()


if __name__ == "__main__":
    main()
